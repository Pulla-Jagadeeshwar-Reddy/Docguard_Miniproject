"""Build printable PDF and DOCX copies of the synthetic policy Markdown files."""

from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "demo-policies"
OUTPUT = SOURCE / "files"
PDF_DIR = OUTPUT / "pdf"
DOCX_DIR = OUTPUT / "docx"


def policy_slug(title: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")


def inline_markup(text: str) -> str:
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)


def generate_pdf(title: str, lines: list[str]) -> None:
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "PolicyTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=15, leading=18, textColor=colors.HexColor("#18324a"),
        alignment=TA_CENTER, spaceAfter=8,
    )
    h2_style = ParagraphStyle(
        "PolicyHeading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=10, leading=12, textColor=colors.HexColor("#244b64"),
        spaceBefore=6, spaceAfter=2, keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "PolicyBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=8.5, leading=10.4, textColor=colors.HexColor("#202a33"),
        spaceAfter=4,
    )
    story = [Paragraph(title, title_style), Paragraph(
        "Synthetic educational example | Not clinical, legal, or operational guidance",
        ParagraphStyle("Notice", parent=body_style, alignment=TA_CENTER, textColor=colors.HexColor("#65717d"), fontSize=7.5, leading=9),
    ), Spacer(1, 5)]
    for line in lines:
        if not line.strip():
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            story.append(Paragraph(inline_markup(line[3:]), h2_style))
        else:
            story.append(Paragraph(inline_markup(line), body_style))

    target = PDF_DIR / f"{policy_slug(title)}.pdf"
    document = SimpleDocTemplate(
        str(target), pagesize=LETTER, rightMargin=0.62 * inch,
        leftMargin=0.62 * inch, topMargin=0.52 * inch, bottomMargin=0.52 * inch,
        title=title, author="DocuGuard synthetic demo",
    )

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#65717d"))
        canvas.drawString(0.62 * inch, 0.3 * inch, "DocuGuard synthetic policy")
        canvas.drawRightString(LETTER[0] - 0.62 * inch, 0.3 * inch, f"Page {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)


def generate_docx(title: str, lines: list[str]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(32, 42, 51)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    title_style = document.styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(17)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(24, 50, 74)
    title_style.paragraph_format.space_after = Pt(5)

    heading_style = document.styles["Heading 2"]
    heading_style.font.name = "Aptos"
    heading_style.font.size = Pt(10)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(36, 75, 100)
    heading_style.paragraph_format.space_before = Pt(5)
    heading_style.paragraph_format.space_after = Pt(2)

    document.add_paragraph(title, style="Title")
    notice = document.add_paragraph("Synthetic educational example | Not clinical, legal, or operational guidance")
    notice.paragraph_format.space_after = Pt(6)
    for run in notice.runs:
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(101, 113, 125)

    for line in lines:
        if not line.strip() or line.startswith("# "):
            continue
        if line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 2")
            continue
        paragraph = document.add_paragraph()
        for part in re.split(r"(\*\*.+?\*\*)", line):
            if part.startswith("**") and part.endswith("**"):
                paragraph.add_run(part[2:-2]).bold = True
            else:
                paragraph.add_run(part)

    footer = section.footer.paragraphs[0]
    footer.alignment = TA_CENTER
    footer.add_run("DocuGuard synthetic policy").font.size = Pt(7)

    target = DOCX_DIR / f"{policy_slug(title)}.docx"
    document.core_properties.title = title
    document.core_properties.subject = "Synthetic DBMS course demonstration policy"
    document.save(target)


def main() -> None:
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    sources = sorted(SOURCE.glob("*.md"))
    if len(sources) != 12:
        raise SystemExit(f"Expected exactly 12 policy Markdown sources; found {len(sources)}")

    for source in sources:
        text = source.read_text(encoding="utf-8")
        title_line = next((line for line in text.splitlines() if line.startswith("# ")), None)
        if not title_line:
            raise SystemExit(f"Missing title in {source.name}")
        title = title_line[2:].strip()
        lines = text.splitlines()
        generate_pdf(title, lines)
        generate_docx(title, lines)
    print(f"Generated {len(sources)} synthetic policy PDFs and {len(sources)} DOCX files.")


if __name__ == "__main__":
    main()
